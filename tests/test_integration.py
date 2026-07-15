"""Integration tests that run Behave with the registered formatters.

These tests execute ``behave`` as a subprocess against the example project
in ``examples/behave_project`` to verify that the formatters work end-to-end
with a real Behave run.
"""

from __future__ import annotations

import csv
import json
import shutil
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook
from odf.opendocument import load as load_ods
from odf.table import Table

EXAMPLE_PROJECT = Path(__file__).resolve().parent.parent / "examples" / "behave_project"

pytestmark = pytest.mark.skipif(
    shutil.which("behave") is None,
    reason="behave CLI not installed",
)


def _run_behave(
    tmp_path: Path,
    formatter: str,
    output_file: str,
    extra_args: list[str] | None = None,
) -> subprocess.CompletedProcess[str]:
    """Run behave with a formatter against a copy of the example project.

    Args:
        tmp_path: Temporary directory for the run.
        formatter: Formatter name (e.g. ``csv-modern``).
        output_file: Output file name for the report.
        extra_args: Additional userdata key=value pairs.

    Returns:
        The completed process result.
    """
    project_dir = tmp_path / "project"
    shutil.copytree(EXAMPLE_PROJECT, project_dir)

    output_path = tmp_path / output_file
    history_path = tmp_path / "history.json"

    cmd = [
        sys.executable,
        "-m",
        "behave",
        "-f",
        formatter,
        "-o",
        str(output_path),
        "--no-color",
        f"-D report_history_path={history_path}",
    ]
    if extra_args:
        cmd.extend(extra_args)
    cmd.append("features")

    return subprocess.run(
        cmd,
        capture_output=True,
        text=True,
        timeout=60,
        cwd=str(project_dir),
    )


# ---------------------------------------------------------------------------
# CSV formatter integration
# ---------------------------------------------------------------------------


class TestCSVIntegration:
    """CSV formatter produces valid output from a real Behave run."""

    def test_csv_report_has_rows(self, tmp_path: Path) -> None:
        """CSV report contains expected scenario rows from the example features."""
        result = _run_behave(tmp_path, "csv-modern", "report.csv")
        assert result.returncode in (0, 1), f"behave failed: {result.stderr}"

        output = tmp_path / "report.csv"
        assert output.exists(), "CSV report was not created"

        with output.open(encoding="utf-8") as f:
            reader = csv.DictReader(f)
            rows = list(reader)

        assert len(rows) >= 4, f"Expected at least 4 rows, got {len(rows)}"
        statuses = {r["status"] for r in rows}
        assert "passed" in statuses


# ---------------------------------------------------------------------------
# XLSX formatter integration
# ---------------------------------------------------------------------------


class TestXLSXIntegration:
    """XLSX formatter produces valid output from a real Behave run."""

    def test_xlsx_report_valid(self, tmp_path: Path) -> None:
        """XLSX report is a valid workbook with Summary and Details sheets."""
        result = _run_behave(tmp_path, "xlsx-modern", "report.xlsx")
        assert result.returncode in (0, 1), f"behave failed: {result.stderr}"

        output = tmp_path / "report.xlsx"
        assert output.exists(), "XLSX report was not created"

        wb = load_workbook(output)
        assert "Summary" in wb.sheetnames
        assert "Details" in wb.sheetnames
        assert "Trends" in wb.sheetnames
        wb.close()

    def test_xlsx_trends_populated(self, tmp_path: Path) -> None:
        """Trends sheet has one entry after a single run."""
        result = _run_behave(tmp_path, "xlsx-modern", "report.xlsx")
        assert result.returncode in (0, 1)

        output = tmp_path / "report.xlsx"
        wb = load_workbook(output)
        ws_trends = wb["Trends"]
        assert ws_trends.max_row >= 2, "Trends sheet should have header + 1 entry"
        wb.close()


# ---------------------------------------------------------------------------
# ODS formatter integration
# ---------------------------------------------------------------------------


class TestODSIntegration:
    """ODS formatter produces valid output from a real Behave run."""

    def test_ods_report_valid(self, tmp_path: Path) -> None:
        """ODS report is a valid document with Summary and Details tables."""
        result = _run_behave(tmp_path, "ods-modern", "report.ods")
        assert result.returncode in (0, 1), f"behave failed: {result.stderr}"

        output = tmp_path / "report.ods"
        assert output.exists(), "ODS report was not created"

        doc = load_ods(str(output))
        table_names = {
            t.getAttribute("name") for t in doc.spreadsheet.getElementsByType(Table)
        }
        assert "Summary" in table_names
        assert "Details" in table_names
        assert "Trends" in table_names


# ---------------------------------------------------------------------------
# Multi-run history accumulation
# ---------------------------------------------------------------------------


class TestHistoryAccumulation:
    """Running twice produces two trend entries."""

    def test_two_runs_two_trends(self, tmp_path: Path) -> None:
        """Two consecutive XLSX runs produce two trend entries."""
        project_dir = tmp_path / "project"
        shutil.copytree(EXAMPLE_PROJECT, project_dir)

        history_path = tmp_path / "history.json"
        output_path = tmp_path / "report.xlsx"

        for _ in range(2):
            cmd = [
                sys.executable,
                "-m",
                "behave",
                "-f",
                "xlsx-modern",
                "-o",
                str(output_path),
                "--no-color",
                f"-D report_history_path={history_path}",
                "features",
            ]
            result = subprocess.run(
                cmd, capture_output=True, text=True, timeout=60, cwd=str(project_dir)
            )
            assert result.returncode in (0, 1)

        wb = load_workbook(output_path)
        ws_trends = wb["Trends"]
        assert ws_trends.max_row >= 3, "Trends should have header + 2 entries"
        wb.close()

        data = json.loads(history_path.read_text(encoding="utf-8"))
        assert len(data) == 2
