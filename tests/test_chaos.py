"""Chaos and edge case tests for robustness validation.

These tests go beyond normal coverage by stressing the code with unusual
inputs, boundary conditions, concurrent-like sequences, and invalid data
that would not occur in typical usage but must be handled gracefully.
"""

from __future__ import annotations

import csv
import json
import sys
from io import StringIO
from pathlib import Path
from types import SimpleNamespace

import pytest

from behave_modern_sheets_report.collector import Collector
from behave_modern_sheets_report.csv_formatter import CSVFormatter
from behave_modern_sheets_report.csv_writer import COLUMN_MAP, CSVWriter
from behave_modern_sheets_report.history import History
from behave_modern_sheets_report.models import (
    FeatureSummary,
    HistoryEntry,
    RunSummary,
    ScenarioResult,
)
from behave_modern_sheets_report.ods_formatter import ODSFormatter
from behave_modern_sheets_report.utils import (
    DEFAULT_COLUMNS,
    STATUS_FAILED,
    STATUS_PASSED,
    STATUS_SKIPPED,
    STATUS_UNDEFINED,
    STATUS_UNTESTED,
    format_duration,
    generate_id,
    normalize_status,
    parse_bool,
    parse_columns,
    parse_delimiter,
    safe_str,
    safe_tags,
)
from behave_modern_sheets_report.xlsx_formatter import XLSXFormatter
from behave_modern_sheets_report.xlsx_writer import XLSXWriter
from tests._helpers import (
    StreamOpener as _StreamOpener,
)
from tests._helpers import (
    make_feature_obj as _make_feature_obj,
)
from tests._helpers import (
    make_run_summary as _make_run_summary,
)
from tests._helpers import (
    make_scenario_obj as _make_scenario_obj,
)
from tests._helpers import (
    make_scenario_result as _make_scenario_result,
)
from tests._helpers import (
    make_step_obj as _make_step_obj,
)

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


# ---------------------------------------------------------------------------
# Collector chaos tests
# ---------------------------------------------------------------------------


class TestCollectorChaos:
    """Collector under chaotic conditions."""

    def test_start_feature_without_ending_previous(self) -> None:
        c = Collector()
        c.start_feature(_make_feature_obj("F1"))
        c.start_feature(_make_feature_obj("F2"))
        rs = c.finalize()
        assert len(rs.features) == 2
        assert rs.features[0].feature_name == "F1"
        assert rs.features[1].feature_name == "F2"

    def test_finalize_called_twice(self) -> None:
        c = Collector()
        c.start_feature(_make_feature_obj())
        c.start_scenario(_make_scenario_obj())
        c.start_step(_make_step_obj("passed"))
        c.end_step(_make_step_obj("passed"))
        c.end_scenario()
        c.end_feature()
        rs1 = c.finalize()
        rs2 = c.finalize()
        assert rs1.total_scenarios == rs2.total_scenarios
        assert rs1.passed == rs2.passed

    def test_step_with_non_string_status(self) -> None:
        step = SimpleNamespace(name="step", status=42, error=None)
        c = Collector()
        c.start_feature(_make_feature_obj())
        c.start_scenario(_make_scenario_obj())
        c.start_step(step)
        c.end_step(step)
        c.end_scenario()
        c.end_feature()
        rs = c.finalize()
        assert rs.scenarios[0].status == STATUS_UNDEFINED

    def test_feature_with_very_long_name(self) -> None:
        long_name = "A" * 600
        c = Collector()
        c.start_feature(SimpleNamespace(name=long_name, filename="features/x.feature"))
        c.end_feature()
        rs = c.finalize()
        assert len(rs.features[0].feature_name) == 503
        assert rs.features[0].feature_name.endswith("...")

    def test_feature_name_falls_back_to_filename(self) -> None:
        feature = SimpleNamespace(name="", filename="features/checkout.feature")
        c = Collector()
        c.start_feature(feature)
        c.end_feature()
        rs = c.finalize()
        assert rs.features[0].feature_name == "features/checkout.feature"

    def test_scenario_with_no_name(self) -> None:
        scenario = SimpleNamespace(
            name="",
            tags=[],
            location=SimpleNamespace(filename="features/x.feature", line=1),
            rule=None,
            is_outline=False,
        )
        c = Collector()
        c.start_feature(_make_feature_obj())
        c.start_scenario(scenario)
        c.start_step(_make_step_obj("passed"))
        c.end_step(_make_step_obj("passed"))
        c.end_scenario()
        c.end_feature()
        rs = c.finalize()
        assert rs.scenarios[0].scenario_name == ""

    def test_multiple_failed_steps_one_scenario(self) -> None:
        c = Collector()
        c.start_feature(_make_feature_obj())
        c.start_scenario(_make_scenario_obj())
        c.start_step(_make_step_obj("failed", error=ValueError("e1")))
        c.end_step(_make_step_obj("failed", error=ValueError("e1")))
        c.start_step(_make_step_obj("failed", error=ValueError("e2")))
        c.end_step(_make_step_obj("failed", error=ValueError("e2")))
        c.end_scenario()
        c.end_feature()
        rs = c.finalize()
        assert rs.scenarios[0].status == STATUS_FAILED
        assert rs.scenarios[0].failed_steps == 2
        assert "e1" in rs.scenarios[0].error_message

    def test_scenario_with_zero_steps_is_passed(self) -> None:
        c = Collector()
        c.start_feature(_make_feature_obj())
        c.start_scenario(_make_scenario_obj())
        c.end_scenario()
        c.end_feature()
        rs = c.finalize()
        assert rs.scenarios[0].status == STATUS_PASSED
        assert rs.scenarios[0].step_count == 0

    def test_step_with_exception_fallback_attribute(self) -> None:
        step = SimpleNamespace(
            name="step",
            status="failed",
            error=None,
            exception=ValueError("fallback"),
        )
        c = Collector()
        c.start_feature(_make_feature_obj())
        c.start_scenario(_make_scenario_obj())
        c.start_step(step)
        c.end_step(step)
        c.end_scenario()
        c.end_feature()
        rs = c.finalize()
        assert "fallback" in rs.scenarios[0].error_message
        assert rs.scenarios[0].error_type == "ValueError"

    def test_scenario_with_line_none(self) -> None:
        scenario = SimpleNamespace(
            name="S",
            tags=[],
            location=SimpleNamespace(filename="features/x.feature", line=None),
            rule=None,
            is_outline=False,
        )
        c = Collector()
        c.start_feature(_make_feature_obj())
        c.start_scenario(scenario)
        c.start_step(_make_step_obj("passed"))
        c.end_step(_make_step_obj("passed"))
        c.end_scenario()
        c.end_feature()
        rs = c.finalize()
        assert rs.scenarios[0].line == 0


# ---------------------------------------------------------------------------
# History chaos tests
# ---------------------------------------------------------------------------


class TestHistoryChaos:
    """History under chaotic conditions."""

    def test_deserialize_string_numeric_values(self, tmp_path: Path) -> None:
        path = tmp_path / "h.json"
        path.write_text(
            json.dumps([{"run_id": "r1", "passed": "5", "failed": "1", "pass_rate": "83.3"}]),
            encoding="utf-8",
        )
        hist = History(path=path)
        loaded = hist.load()
        assert len(loaded) == 1
        assert loaded[0].passed == 5
        assert loaded[0].failed == 1
        assert loaded[0].pass_rate == 83.3

    def test_deserialize_non_numeric_string_skips_entry(self, tmp_path: Path) -> None:
        path = tmp_path / "h.json"
        path.write_text(
            json.dumps([{"run_id": "r1", "passed": "not_a_number"}]),
            encoding="utf-8",
        )
        hist = History(path=path)
        assert hist.load() == []

    def test_deserialize_with_extra_unknown_fields(self, tmp_path: Path) -> None:
        path = tmp_path / "h.json"
        path.write_text(
            json.dumps([{"run_id": "r1", "passed": 5, "unknown_field": "ignored", "extra": 42}]),
            encoding="utf-8",
        )
        hist = History(path=path)
        loaded = hist.load()
        assert len(loaded) == 1
        assert loaded[0].run_id == "r1"
        assert loaded[0].passed == 5

    def test_append_exactly_max_entries_no_truncation(self, tmp_path: Path) -> None:
        hist = History(path=tmp_path / "h.json", max_entries=3)
        for _i in range(3):
            hist.append(_make_run_summary())
        loaded = hist.load()
        assert len(loaded) == 3

    def test_clear_then_append(self, tmp_path: Path) -> None:
        hist = History(path=tmp_path / "h.json")
        hist.append(_make_run_summary())
        hist.append(_make_run_summary())
        assert len(hist.load()) == 2
        hist.clear()
        assert hist.load() == []
        hist.append(_make_run_summary())
        assert len(hist.load()) == 1

    def test_serialize_produces_valid_json(self, tmp_path: Path) -> None:
        hist = History(path=tmp_path / "h.json")
        entries = [HistoryEntry(run_id=f"r{i}", passed=i) for i in range(3)]
        raw = hist._serialize(entries)
        parsed = json.loads(raw)
        assert isinstance(parsed, list)
        assert len(parsed) == 3
        assert parsed[0]["run_id"] == "r0"

    def test_nested_directory_created_automatically(self, tmp_path: Path) -> None:
        deep = tmp_path / "a" / "b" / "c" / "d" / "h.json"
        hist = History(path=deep)
        hist.append(_make_run_summary())
        assert deep.exists()

    def test_deserialize_empty_dict_entry(self, tmp_path: Path) -> None:
        path = tmp_path / "h.json"
        path.write_text(json.dumps([{}]), encoding="utf-8")
        hist = History(path=path)
        loaded = hist.load()
        assert len(loaded) == 1
        assert loaded[0].run_id == ""
        assert loaded[0].passed == 0

    def test_deserialize_mixed_valid_invalid_entries(self, tmp_path: Path) -> None:
        path = tmp_path / "h.json"
        path.write_text(
            json.dumps([42, {"run_id": "valid", "passed": 1}, "string", {"run_id": "valid2"}]),
            encoding="utf-8",
        )
        hist = History(path=path)
        loaded = hist.load()
        assert len(loaded) == 2
        assert loaded[0].run_id == "valid"
        assert loaded[1].run_id == "valid2"

    def test_load_valid_json_array_with_empty_list(self, tmp_path: Path) -> None:
        path = tmp_path / "h.json"
        path.write_text("[]", encoding="utf-8")
        hist = History(path=path)
        assert hist.load() == []

    def test_max_entries_one_hundred_default(self, tmp_path: Path) -> None:
        hist = History(path=tmp_path / "h.json")
        assert hist._max_entries == 100

    def test_append_returns_truncated_list(self, tmp_path: Path) -> None:
        hist = History(path=tmp_path / "h.json", max_entries=2)
        hist.append(_make_run_summary())
        entries = hist.append(_make_run_summary())
        assert len(entries) == 2
        entries_after = hist.append(_make_run_summary())
        assert len(entries_after) == 2


# ---------------------------------------------------------------------------
# Utils chaos tests
# ---------------------------------------------------------------------------


class TestUtilsChaos:
    """Utils under chaotic conditions."""

    def test_safe_str_float(self) -> None:
        assert safe_str(3.14) == "3.14"

    def test_safe_str_bool_true(self) -> None:
        assert safe_str(True) == "True"

    def test_safe_str_bool_false(self) -> None:
        assert safe_str(False) == "False"

    def test_safe_str_list(self) -> None:
        assert safe_str([1, 2]) == "[1, 2]"

    def test_safe_str_dict(self) -> None:
        result = safe_str({"a": 1})
        assert "a" in result
        assert "1" in result

    def test_safe_str_501_chars_truncated(self) -> None:
        result = safe_str("a" * 501)
        assert len(result) == 503
        assert result.endswith("...")

    def test_safe_tags_list_of_ints(self) -> None:
        assert safe_tags([1, 2, 3]) == ["1", "2", "3"]

    def test_safe_tags_list_with_none(self) -> None:
        assert safe_tags([None, "a", ""]) == ["None", "a"]

    def test_safe_tags_dict_returns_empty(self) -> None:
        assert safe_tags({"a": 1}) == []

    def test_safe_tags_tuple_returns_empty(self) -> None:
        assert safe_tags(("a", "b")) == []

    def test_format_duration_very_large(self) -> None:
        result = format_duration(999999.999)
        assert result == "999999.999s"

    def test_format_duration_just_below_1ms(self) -> None:
        assert format_duration(0.0009) == "0s"

    def test_format_duration_just_above_1ms(self) -> None:
        assert format_duration(0.0011) == "1ms"

    def test_format_duration_just_below_1s(self) -> None:
        assert format_duration(0.999) == "999ms"

    def test_format_duration_just_at_1s(self) -> None:
        assert format_duration(1.0) == "1.000s"

    def test_generate_id_empty_prefix(self) -> None:
        result = generate_id("")
        assert result.startswith("_")
        assert len(result) == 13

    def test_generate_id_prefix_with_special_chars(self) -> None:
        result = generate_id("feat-1.0")
        assert result.startswith("feat-1.0_")

    def test_parse_columns_all_empty_entries(self) -> None:
        assert parse_columns(",,,") == DEFAULT_COLUMNS

    def test_parse_columns_single_comma(self) -> None:
        assert parse_columns(",") == DEFAULT_COLUMNS

    def test_parse_bool_empty_string(self) -> None:
        assert parse_bool("") is False

    def test_parse_bool_whitespace_only(self) -> None:
        assert parse_bool("   ") is False

    def test_parse_delimiter_empty_string(self) -> None:
        assert parse_delimiter("") == ","

    def test_parse_delimiter_whitespace_only(self) -> None:
        assert parse_delimiter("   ") == ","

    def test_normalize_status_with_whitespace(self) -> None:
        assert normalize_status("  passed  ") == STATUS_PASSED
        assert normalize_status("  failed  ") == STATUS_FAILED

    def test_normalize_status_with_mixed_case(self) -> None:
        assert normalize_status("PASSED") == STATUS_PASSED
        assert normalize_status("FaIlEd") == STATUS_FAILED
        assert normalize_status("SkIpPeD") == STATUS_SKIPPED

    def test_normalize_status_with_uppercase_undefined(self) -> None:
        assert normalize_status("UNDEFINED") == STATUS_UNDEFINED

    def test_parse_columns_returns_copy_of_default(self) -> None:
        result = parse_columns(None)
        assert result == DEFAULT_COLUMNS
        result.append("custom")
        assert "custom" not in DEFAULT_COLUMNS


# ---------------------------------------------------------------------------
# CSV Writer chaos tests
# ---------------------------------------------------------------------------


class TestCSVWriterChaos:
    """CSV Writer under chaotic conditions."""

    def test_special_characters_in_scenario_name(self) -> None:
        scenario = _make_scenario_result()
        scenario.scenario_name = 'Test "with quotes", commas, and\nnewlines'
        run = _make_run_summary(scenarios=[scenario])
        stream = StringIO()
        CSVWriter.write(run, stream)
        stream.seek(0)
        reader = csv.DictReader(stream)
        rows = list(reader)
        assert len(rows) == 1
        assert "quotes" in rows[0]["scenario"]
        assert "commas" in rows[0]["scenario"]

    def test_all_columns_selected(self) -> None:
        scenario = _make_scenario_result(
            status=STATUS_FAILED,
            tags=["a", "b"],
            error_message="boom",
            error_type="ValueError",
            is_outline=True,
        )
        run = _make_run_summary(scenarios=[scenario])
        stream = StringIO()
        all_cols = list(COLUMN_MAP.keys())
        CSVWriter.write(run, stream, columns=all_cols)
        stream.seek(0)
        reader = csv.DictReader(stream)
        rows = list(reader)
        assert len(rows) == 1
        for col in all_cols:
            assert col in rows[0]

    def test_unknown_column_raises_keyerror(self) -> None:
        run = _make_run_summary(scenarios=[_make_scenario_result()])
        stream = StringIO()
        with pytest.raises(KeyError):
            CSVWriter.write(run, stream, columns=["nonexistent_column"])

    def test_large_number_of_scenarios(self) -> None:
        scenarios = [
            ScenarioResult(feature_name="F", scenario_name=f"S{i}", status=STATUS_PASSED)
            for i in range(500)
        ]
        run = RunSummary(scenarios=scenarios, total_scenarios=500, passed=500, pass_rate=100.0)
        stream = StringIO()
        CSVWriter.write(run, stream)
        stream.seek(0)
        reader = csv.DictReader(stream)
        rows = list(reader)
        assert len(rows) == 500

    def test_unicode_in_scenario_name(self) -> None:
        scenario = _make_scenario_result()
        scenario.scenario_name = "Test with unicode: \u00e9\u00e8\u00ea\u00eb \u4e2d\u6587"
        run = _make_run_summary(scenarios=[scenario])
        stream = StringIO()
        CSVWriter.write(run, stream)
        stream.seek(0)
        reader = csv.DictReader(stream)
        rows = list(reader)
        assert "\u00e9" in rows[0]["scenario"]
        assert "\u4e2d" in rows[0]["scenario"]

    def test_empty_run_summary(self) -> None:
        run = RunSummary()
        stream = StringIO()
        CSVWriter.write(run, stream)
        stream.seek(0)
        reader = csv.reader(stream)
        header = next(reader)
        assert header == DEFAULT_COLUMNS
        assert list(reader) == []

    def test_all_failed_scenarios(self) -> None:
        scenarios = [
            _make_scenario_result(status=STATUS_FAILED, error_message=f"err{i}") for i in range(5)
        ]
        run = _make_run_summary(scenarios=scenarios)
        stream = StringIO()
        CSVWriter.write(run, stream, only_failed=True)
        stream.seek(0)
        reader = csv.DictReader(stream)
        rows = list(reader)
        assert len(rows) == 5
        for row in rows:
            assert row["status"] == "failed"

    def test_only_failed_with_no_failed_scenarios(self) -> None:
        scenarios = [
            _make_scenario_result(status=STATUS_PASSED),
            _make_scenario_result(status=STATUS_SKIPPED),
        ]
        run = _make_run_summary(scenarios=scenarios)
        stream = StringIO()
        CSVWriter.write(run, stream, only_failed=True)
        stream.seek(0)
        reader = csv.DictReader(stream)
        rows = list(reader)
        assert len(rows) == 0

    def test_traceback_column(self) -> None:
        scenario = _make_scenario_result(status=STATUS_FAILED, error_message="boom")
        scenario.traceback = "Traceback (most recent call last):\n  File ..."
        run = _make_run_summary(scenarios=[scenario])
        stream = StringIO()
        CSVWriter.write(run, stream, columns=["feature", "traceback"])
        stream.seek(0)
        reader = csv.DictReader(stream)
        rows = list(reader)
        assert "Traceback" in rows[0]["traceback"]

    def test_all_step_columns(self) -> None:
        scenario = _make_scenario_result(status=STATUS_FAILED)
        scenario.step_count = 10
        scenario.passed_steps = 7
        scenario.failed_steps = 2
        scenario.skipped_steps = 1
        run = _make_run_summary(scenarios=[scenario])
        stream = StringIO()
        CSVWriter.write(
            run,
            stream,
            columns=["steps", "passed_steps", "failed_steps", "skipped_steps"],
        )
        stream.seek(0)
        reader = csv.DictReader(stream)
        rows = list(reader)
        assert rows[0]["steps"] == "10"
        assert rows[0]["passed_steps"] == "7"
        assert rows[0]["failed_steps"] == "2"
        assert rows[0]["skipped_steps"] == "1"


# ---------------------------------------------------------------------------
# XLSX Writer chaos tests
# ---------------------------------------------------------------------------


class TestXLSXWriterChaos:
    """XLSX Writer under chaotic conditions."""

    def test_all_scenarios_failed(self, tmp_path: Path) -> None:
        from openpyxl import load_workbook

        scenarios = [
            _make_scenario_result(status=STATUS_FAILED, error_message=f"err{i}") for i in range(3)
        ]
        run = _make_run_summary(scenarios=scenarios)
        path = tmp_path / "report.xlsx"
        XLSXWriter.write(run, path)
        wb = load_workbook(path)
        ws_failures = wb["Failures"]
        assert ws_failures.max_row == 4
        wb.close()

    def test_all_scenarios_skipped(self, tmp_path: Path) -> None:
        from openpyxl import load_workbook

        scenarios = [_make_scenario_result(status=STATUS_SKIPPED) for _ in range(3)]
        run = _make_run_summary(scenarios=scenarios)
        path = tmp_path / "report.xlsx"
        XLSXWriter.write(run, path)
        wb = load_workbook(path)
        ws_failures = wb["Failures"]
        assert ws_failures.max_row == 1
        wb.close()

    def test_all_columns_in_details(self, tmp_path: Path) -> None:
        from openpyxl import load_workbook

        scenario = _make_scenario_result(
            status=STATUS_FAILED,
            tags=["a", "b"],
            error_message="boom",
            error_type="ValueError",
            is_outline=True,
        )
        run = _make_run_summary(scenarios=[scenario])
        path = tmp_path / "report.xlsx"
        all_cols = list(COLUMN_MAP.keys())
        XLSXWriter.write(run, path, columns=all_cols)
        wb = load_workbook(path)
        ws = wb["Details"]
        header = [cell.value for cell in ws[1]]
        assert header == all_cols
        wb.close()

    def test_many_trends_entries(self, tmp_path: Path) -> None:
        from openpyxl import load_workbook

        trends = [HistoryEntry(run_id=f"r{i}", pass_rate=float(i), passed=i) for i in range(50)]
        run = _make_run_summary()
        path = tmp_path / "report.xlsx"
        XLSXWriter.write(run, path, trends=trends)
        wb = load_workbook(path)
        ws = wb["Trends"]
        assert ws.max_row == 51
        wb.close()

    def test_undefined_status_no_fill(self, tmp_path: Path) -> None:
        from openpyxl import load_workbook

        scenario = _make_scenario_result(status=STATUS_UNDEFINED)
        run = _make_run_summary(scenarios=[scenario])
        path = tmp_path / "report.xlsx"
        XLSXWriter.write(run, path)
        wb = load_workbook(path)
        ws = wb["Details"]
        cell = ws.cell(row=2, column=3)
        assert cell.fill.start_color.rgb in (None, "00000000")
        wb.close()

    def test_untested_status_no_fill(self, tmp_path: Path) -> None:
        from openpyxl import load_workbook

        scenario = _make_scenario_result(status=STATUS_UNTESTED)
        run = _make_run_summary(scenarios=[scenario])
        path = tmp_path / "report.xlsx"
        XLSXWriter.write(run, path)
        wb = load_workbook(path)
        ws = wb["Details"]
        cell = ws.cell(row=2, column=3)
        assert cell.fill.start_color.rgb in (None, "00000000")
        wb.close()

    def test_empty_features_summary(self, tmp_path: Path) -> None:
        from openpyxl import load_workbook

        run = RunSummary()
        path = tmp_path / "report.xlsx"
        XLSXWriter.write(run, path)
        wb = load_workbook(path)
        ws = wb["Summary"]
        assert ws.max_row == 1
        wb.close()

    def test_output_to_nested_directory(self, tmp_path: Path) -> None:
        from openpyxl import load_workbook

        run = _make_run_summary(scenarios=[_make_scenario_result()])
        path = tmp_path / "a" / "b" / "c" / "report.xlsx"
        path.parent.mkdir(parents=True, exist_ok=True)
        XLSXWriter.write(run, path)
        assert path.exists()
        wb = load_workbook(path)
        assert "Summary" in wb.sheetnames
        wb.close()


# ---------------------------------------------------------------------------
# ODS Writer chaos tests
# ---------------------------------------------------------------------------


class TestODSWriterChaos:
    """ODS Writer under chaotic conditions."""

    def test_all_scenarios_failed(self, tmp_path: Path) -> None:
        from odf.opendocument import load
        from odf.table import Table

        from behave_modern_sheets_report.ods_writer import ODSWriter

        scenarios = [
            _make_scenario_result(status=STATUS_FAILED, error_message=f"err{i}") for i in range(3)
        ]
        run = _make_run_summary(scenarios=scenarios)
        path = tmp_path / "report.ods"
        ODSWriter.write(run, path)
        doc = load(str(path))
        tables = {t.getAttribute("name") for t in doc.spreadsheet.getElementsByType(Table)}
        assert "Failures" in tables

    def test_all_columns_in_details(self, tmp_path: Path) -> None:
        from odf.opendocument import load
        from odf.table import Table

        from behave_modern_sheets_report.ods_writer import ODSWriter

        scenario = _make_scenario_result(
            status=STATUS_FAILED,
            tags=["a", "b"],
            error_message="boom",
            error_type="ValueError",
            is_outline=True,
        )
        run = _make_run_summary(scenarios=[scenario])
        path = tmp_path / "report.ods"
        all_cols = list(COLUMN_MAP.keys())
        ODSWriter.write(run, path, columns=all_cols)
        doc = load(str(path))
        tables = {t.getAttribute("name") for t in doc.spreadsheet.getElementsByType(Table)}
        assert "Details" in tables

    def test_many_trends_entries(self, tmp_path: Path) -> None:
        from odf.opendocument import load
        from odf.table import Table

        from behave_modern_sheets_report.ods_writer import ODSWriter

        trends = [HistoryEntry(run_id=f"r{i}", pass_rate=float(i)) for i in range(30)]
        run = _make_run_summary()
        path = tmp_path / "report.ods"
        ODSWriter.write(run, path, trends=trends)
        doc = load(str(path))
        tables = {t.getAttribute("name") for t in doc.spreadsheet.getElementsByType(Table)}
        assert "Trends" in tables

    def test_output_to_nested_directory(self, tmp_path: Path) -> None:
        from odf.opendocument import load

        from behave_modern_sheets_report.ods_writer import ODSWriter

        run = _make_run_summary(scenarios=[_make_scenario_result()])
        path = tmp_path / "x" / "y" / "report.ods"
        path.parent.mkdir(parents=True, exist_ok=True)
        ODSWriter.write(run, path)
        assert path.exists()
        doc = load(str(path))
        assert doc.spreadsheet is not None


# ---------------------------------------------------------------------------
# CSV Formatter chaos tests
# ---------------------------------------------------------------------------


class TestCSVFormatterChaos:
    """CSV Formatter under chaotic conditions."""

    def test_double_close_no_crash(self) -> None:
        stream = StringIO()
        opener = _StreamOpener(stream)
        config = SimpleNamespace(userdata={})
        fmt = CSVFormatter(opener, config)
        feature = _make_feature_obj("Login")
        scenario = _make_scenario_obj("S1")
        step = _make_step_obj("passed")
        fmt.uri("features/login.feature")
        fmt.feature(feature)
        fmt.scenario(scenario)
        fmt.step(step)
        fmt.result(step)
        fmt.close()
        fmt.close()
        stream.seek(0)
        reader = csv.reader(stream)
        lines = list(reader)
        assert len(lines) == 2

    def test_result_without_step_start(self) -> None:
        stream = StringIO()
        opener = _StreamOpener(stream)
        config = SimpleNamespace(userdata={})
        fmt = CSVFormatter(opener, config)
        feature = _make_feature_obj()
        scenario = _make_scenario_obj()
        fmt.uri("features/login.feature")
        fmt.feature(feature)
        fmt.scenario(scenario)
        fmt.result(_make_step_obj("passed"))
        fmt.close()
        stream.seek(0)
        reader = csv.DictReader(stream)
        rows = list(reader)
        assert len(rows) == 1
        assert rows[0]["status"] == "passed"

    def test_scenario_auto_finalizes_previous(self) -> None:
        stream = StringIO()
        opener = _StreamOpener(stream)
        config = SimpleNamespace(userdata={})
        fmt = CSVFormatter(opener, config)
        feature = _make_feature_obj()
        s1 = _make_scenario_obj("S1")
        s2 = _make_scenario_obj("S2")
        step = _make_step_obj("passed")
        fmt.uri("features/login.feature")
        fmt.feature(feature)
        fmt.scenario(s1)
        fmt.step(step)
        fmt.result(step)
        fmt.scenario(s2)
        fmt.step(step)
        fmt.result(step)
        fmt.close()
        stream.seek(0)
        reader = csv.DictReader(stream)
        rows = list(reader)
        assert len(rows) == 2
        assert rows[0]["scenario"] == "S1"
        assert rows[1]["scenario"] == "S2"

    def test_multiple_features_no_eof_between(self) -> None:
        stream = StringIO()
        opener = _StreamOpener(stream)
        config = SimpleNamespace(userdata={})
        fmt = CSVFormatter(opener, config)
        step = _make_step_obj("passed")
        fmt.uri("features/f1.feature")
        fmt.feature(_make_feature_obj("F1"))
        fmt.scenario(_make_scenario_obj("S1"))
        fmt.step(step)
        fmt.result(step)
        fmt.uri("features/f2.feature")
        fmt.feature(_make_feature_obj("F2"))
        fmt.scenario(_make_scenario_obj("S2"))
        fmt.step(step)
        fmt.result(step)
        fmt.close()
        stream.seek(0)
        reader = csv.DictReader(stream)
        rows = list(reader)
        assert len(rows) == 2
        assert rows[0]["feature"] == "F1"
        assert rows[1]["feature"] == "F2"

    def test_config_with_non_dict_userdata(self) -> None:
        config = SimpleNamespace(userdata="not a dict")
        fmt = CSVFormatter(None, config)
        assert fmt._columns == DEFAULT_COLUMNS

    def test_stream_opener_open_returns_none(self) -> None:
        opener = SimpleNamespace(open=lambda: None)
        config = SimpleNamespace(userdata={})
        fmt = CSVFormatter(opener, config)
        stream = fmt._resolve_stream()
        assert stream is sys.stdout


# ---------------------------------------------------------------------------
# XLSX Formatter chaos tests
# ---------------------------------------------------------------------------


class TestXLSXFormatterChaos:
    """XLSX Formatter under chaotic conditions."""

    def test_double_close_no_crash(self, tmp_path: Path) -> None:
        from openpyxl import load_workbook

        output = tmp_path / "report.xlsx"
        history = tmp_path / "h.json"
        opener = _StreamOpener(name=str(output))
        config = SimpleNamespace(userdata={"report_history_path": str(history)})
        fmt = XLSXFormatter(opener, config)
        feature = _make_feature_obj()
        scenario = _make_scenario_obj()
        step = _make_step_obj("passed")
        fmt.uri("features/login.feature")
        fmt.feature(feature)
        fmt.scenario(scenario)
        fmt.step(step)
        fmt.result(step)
        fmt.close()
        fmt.close()
        assert output.exists()
        wb = load_workbook(output)
        assert "Summary" in wb.sheetnames
        wb.close()

    def test_multiple_features_no_eof_between(self, tmp_path: Path) -> None:
        from openpyxl import load_workbook

        output = tmp_path / "report.xlsx"
        history = tmp_path / "h.json"
        opener = _StreamOpener(name=str(output))
        config = SimpleNamespace(userdata={"report_history_path": str(history)})
        fmt = XLSXFormatter(opener, config)
        step = _make_step_obj("passed")
        fmt.uri("features/f1.feature")
        fmt.feature(_make_feature_obj("F1"))
        fmt.scenario(_make_scenario_obj("S1"))
        fmt.step(step)
        fmt.result(step)
        fmt.uri("features/f2.feature")
        fmt.feature(_make_feature_obj("F2"))
        fmt.scenario(_make_scenario_obj("S2"))
        fmt.step(step)
        fmt.result(step)
        fmt.close()
        wb = load_workbook(output)
        ws = wb["Summary"]
        assert ws.max_row == 3
        wb.close()

    def test_config_with_non_dict_userdata(self, tmp_path: Path) -> None:
        opener = _StreamOpener(name=str(tmp_path / "report.xlsx"))
        config = SimpleNamespace(userdata="not a dict")
        fmt = XLSXFormatter(opener, config)
        assert fmt._columns == DEFAULT_COLUMNS
        assert fmt._max_history == 100

    def test_stream_opener_without_name_falls_back(self, tmp_path: Path) -> None:
        opener = SimpleNamespace()
        config = SimpleNamespace(userdata={})
        fmt = XLSXFormatter(opener, config)
        assert fmt._resolve_output_path() == "report.xlsx"

    def test_max_history_invalid_string_raises(self, tmp_path: Path) -> None:
        opener = _StreamOpener(name=str(tmp_path / "report.xlsx"))
        config = SimpleNamespace(userdata={"report_max_history": "not_a_number"})
        with pytest.raises(ValueError):
            XLSXFormatter(opener, config)


# ---------------------------------------------------------------------------
# ODS Formatter chaos tests
# ---------------------------------------------------------------------------


class TestODSFormatterChaos:
    """ODS Formatter under chaotic conditions."""

    def test_double_close_no_crash(self, tmp_path: Path) -> None:
        from odf.opendocument import load

        output = tmp_path / "report.ods"
        history = tmp_path / "h.json"
        opener = _StreamOpener(name=str(output))
        config = SimpleNamespace(userdata={"report_history_path": str(history)})
        fmt = ODSFormatter(opener, config)
        feature = _make_feature_obj()
        scenario = _make_scenario_obj()
        step = _make_step_obj("passed")
        fmt.uri("features/login.feature")
        fmt.feature(feature)
        fmt.scenario(scenario)
        fmt.step(step)
        fmt.result(step)
        fmt.close()
        fmt.close()
        assert output.exists()
        doc = load(str(output))
        assert doc.spreadsheet is not None

    def test_multiple_features_no_eof_between(self, tmp_path: Path) -> None:
        from odf.opendocument import load
        from odf.table import Table

        output = tmp_path / "report.ods"
        history = tmp_path / "h.json"
        opener = _StreamOpener(name=str(output))
        config = SimpleNamespace(userdata={"report_history_path": str(history)})
        fmt = ODSFormatter(opener, config)
        step = _make_step_obj("passed")
        fmt.uri("features/f1.feature")
        fmt.feature(_make_feature_obj("F1"))
        fmt.scenario(_make_scenario_obj("S1"))
        fmt.step(step)
        fmt.result(step)
        fmt.uri("features/f2.feature")
        fmt.feature(_make_feature_obj("F2"))
        fmt.scenario(_make_scenario_obj("S2"))
        fmt.step(step)
        fmt.result(step)
        fmt.close()
        doc = load(str(output))
        tables = {t.getAttribute("name") for t in doc.spreadsheet.getElementsByType(Table)}
        assert "Summary" in tables

    def test_config_with_non_dict_userdata(self, tmp_path: Path) -> None:
        opener = _StreamOpener(name=str(tmp_path / "report.ods"))
        config = SimpleNamespace(userdata="not a dict")
        fmt = ODSFormatter(opener, config)
        assert fmt._columns == DEFAULT_COLUMNS
        assert fmt._max_history == 100

    def test_stream_opener_without_name_falls_back(self, tmp_path: Path) -> None:
        opener = SimpleNamespace()
        config = SimpleNamespace(userdata={})
        fmt = ODSFormatter(opener, config)
        assert fmt._resolve_output_path() == "report.ods"


# ---------------------------------------------------------------------------
# Model chaos tests
# ---------------------------------------------------------------------------


class TestModelChaos:
    """Model dataclasses under chaotic conditions."""

    def test_scenario_result_with_all_fields_populated(self) -> None:
        sr = ScenarioResult(
            feature_name="F" * 500,
            scenario_name="S" * 500,
            status="failed",
            duration=999999.999,
            tags=["t"] * 100,
            error_message="e" * 500,
            error_type="CustomError",
            traceback="t" * 1000,
            step_count=999,
            passed_steps=500,
            failed_steps=499,
            skipped_steps=0,
            file="features/very/deep/nested/path.feature",
            line=99999,
            rule="R" * 200,
            is_outline=True,
        )
        assert len(sr.feature_name) == 500
        assert len(sr.tags) == 100
        assert sr.step_count == 999

    def test_feature_summary_with_extreme_values(self) -> None:
        fs = FeatureSummary(
            feature_name="Extreme",
            total_scenarios=1000000,
            passed=999999,
            failed=1,
            skipped=0,
            undefined=0,
            pass_rate=99.9999,
            duration=99999.999,
        )
        assert fs.total_scenarios == 1000000
        assert fs.pass_rate == 99.9999

    def test_history_entry_with_zero_everything(self) -> None:
        he = HistoryEntry(
            run_id="",
            timestamp="",
            total_features=0,
            total_scenarios=0,
            passed=0,
            failed=0,
            skipped=0,
            undefined=0,
            pass_rate=0.0,
            duration=0.0,
        )
        assert he.run_id == ""
        assert he.pass_rate == 0.0

    def test_run_summary_with_large_lists(self) -> None:
        features = [FeatureSummary(feature_name=f"F{i}") for i in range(100)]
        scenarios = [ScenarioResult(scenario_name=f"S{i}") for i in range(1000)]
        rs = RunSummary(features=features, scenarios=scenarios)
        assert len(rs.features) == 100
        assert len(rs.scenarios) == 1000

    def test_scenario_result_equality_with_identical_large_objects(self) -> None:
        sr1 = ScenarioResult(
            feature_name="F",
            tags=["a"] * 50,
            step_count=100,
        )
        sr2 = ScenarioResult(
            feature_name="F",
            tags=["a"] * 50,
            step_count=100,
        )
        assert sr1 == sr2

    def test_feature_summary_inequality_with_different_durations(self) -> None:
        fs1 = FeatureSummary(duration=1.0)
        fs2 = FeatureSummary(duration=2.0)
        assert fs1 != fs2

    def test_run_summary_inequality_with_different_scenarios(self) -> None:
        rs1 = RunSummary(scenarios=[ScenarioResult(scenario_name="A")])
        rs2 = RunSummary(scenarios=[ScenarioResult(scenario_name="B")])
        assert rs1 != rs2
