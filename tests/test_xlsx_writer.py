"""Tests for the XLSX writer module."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from behave_modern_sheets_report.xlsx_writer import XLSXWriter
from tests._helpers import (
    make_feature_summary as _make_feature,
)
from tests._helpers import (
    make_history_entry as _make_history_entry,
)
from tests._helpers import (
    make_run_summary as _make_run_summary,
)
from tests._helpers import (
    make_scenario_result as _make_scenario,
)


class TestValidFile:
    """Generated file is a valid .xlsx readable by openpyxl."""

    def test_file_is_valid(self, tmp_path: Path) -> None:
        """The output file can be loaded by openpyxl.load_workbook."""
        run = _make_run_summary(
            features=[_make_feature()],
            scenarios=[_make_scenario()],
        )
        path = tmp_path / "report.xlsx"
        XLSXWriter.write(run, path)
        wb = load_workbook(path)
        assert "Summary" in wb.sheetnames
        wb.close()


class TestSummarySheet:
    """Summary sheet content."""

    def test_summary_one_row_per_feature(self, tmp_path: Path) -> None:
        """Summary has one data row per feature plus a header."""
        features = [_make_feature("Login"), _make_feature("Auth")]
        run = _make_run_summary(features=features, scenarios=[_make_scenario()])
        path = tmp_path / "report.xlsx"
        XLSXWriter.write(run, path)
        wb = load_workbook(path)
        ws = wb["Summary"]
        assert ws.max_row == 3
        assert ws.cell(row=1, column=1).value == "Feature"
        assert ws.cell(row=2, column=1).value == "Login"
        assert ws.cell(row=3, column=1).value == "Auth"
        wb.close()

    def test_summary_pass_rate_formatted(self, tmp_path: Path) -> None:
        """Pass rate is formatted as a percentage string."""
        features = [_make_feature(pass_rate=85.5)]
        run = _make_run_summary(features=features, scenarios=[_make_scenario()])
        path = tmp_path / "report.xlsx"
        XLSXWriter.write(run, path)
        wb = load_workbook(path)
        ws = wb["Summary"]
        assert ws.cell(row=2, column=7).value == "85.5%"
        wb.close()

    def test_summary_duration_formatted(self, tmp_path: Path) -> None:
        """Duration is formatted with format_duration."""
        features = [_make_feature(duration=1.234)]
        run = _make_run_summary(features=features, scenarios=[_make_scenario()])
        path = tmp_path / "report.xlsx"
        XLSXWriter.write(run, path)
        wb = load_workbook(path)
        ws = wb["Summary"]
        assert ws.cell(row=2, column=8).value == "1.234s"
        wb.close()


class TestDetailsSheet:
    """Details sheet content."""

    def test_details_one_row_per_scenario(self, tmp_path: Path) -> None:
        """Details has one data row per scenario plus a header."""
        scenarios = [
            _make_scenario(scenario_name="S1"),
            _make_scenario(scenario_name="S2"),
        ]
        run = _make_run_summary(
            features=[_make_feature()],
            scenarios=scenarios,
        )
        path = tmp_path / "report.xlsx"
        XLSXWriter.write(run, path)
        wb = load_workbook(path)
        ws = wb["Details"]
        assert ws.max_row == 3
        assert ws.cell(row=2, column=2).value == "S1"
        assert ws.cell(row=3, column=2).value == "S2"
        wb.close()

    def test_details_custom_columns(self, tmp_path: Path) -> None:
        """Custom columns are respected in Details header and rows."""
        run = _make_run_summary(
            features=[_make_feature()],
            scenarios=[_make_scenario()],
        )
        path = tmp_path / "report.xlsx"
        XLSXWriter.write(run, path, columns=["feature", "status"])
        wb = load_workbook(path)
        ws = wb["Details"]
        assert ws.cell(row=1, column=1).value == "feature"
        assert ws.cell(row=1, column=2).value == "status"
        assert ws.cell(row=2, column=1).value == "Login"
        assert ws.cell(row=2, column=2).value == "passed"
        wb.close()


class TestFailuresSheet:
    """Failures sheet content."""

    def test_failures_only_failed(self, tmp_path: Path) -> None:
        """Failures sheet contains only failed scenarios."""
        scenarios = [
            _make_scenario(scenario_name="Pass", status="passed"),
            _make_scenario(
                scenario_name="Fail",
                status="failed",
                error_message="boom",
                error_type="AssertionError",
                traceback="Traceback here",
            ),
        ]
        run = _make_run_summary(
            features=[_make_feature()],
            scenarios=scenarios,
        )
        path = tmp_path / "report.xlsx"
        XLSXWriter.write(run, path)
        wb = load_workbook(path)
        ws = wb["Failures"]
        assert ws.max_row == 2
        assert ws.cell(row=2, column=2).value == "Fail"
        assert ws.cell(row=2, column=3).value == "boom"
        assert ws.cell(row=2, column=4).value == "AssertionError"
        assert ws.cell(row=2, column=5).value == "Traceback here"
        wb.close()

    def test_failures_empty_when_no_failures(self, tmp_path: Path) -> None:
        """Failures sheet has only header when no scenarios failed."""
        run = _make_run_summary(
            features=[_make_feature()],
            scenarios=[_make_scenario(status="passed")],
        )
        path = tmp_path / "report.xlsx"
        XLSXWriter.write(run, path)
        wb = load_workbook(path)
        ws = wb["Failures"]
        assert ws.max_row == 1
        wb.close()


class TestTrendsSheet:
    """Trends sheet content."""

    def test_trends_present(self, tmp_path: Path) -> None:
        """Trends sheet is created when trends is non-empty."""
        trends = [_make_history_entry(), _make_history_entry(timestamp="2025-01-02T12:00:00+00:00")]
        run = _make_run_summary(
            features=[_make_feature()],
            scenarios=[_make_scenario()],
        )
        path = tmp_path / "report.xlsx"
        XLSXWriter.write(run, path, trends=trends)
        wb = load_workbook(path)
        assert "Trends" in wb.sheetnames
        ws = wb["Trends"]
        assert ws.max_row == 3
        assert ws.cell(row=2, column=1).value == "2025-01-01T12:00:00+00:00"
        assert ws.cell(row=2, column=2).value == "85.5%"
        wb.close()

    def test_trends_omitted_when_none(self, tmp_path: Path) -> None:
        """Trends sheet is omitted when trends is None."""
        run = _make_run_summary(
            features=[_make_feature()],
            scenarios=[_make_scenario()],
        )
        path = tmp_path / "report.xlsx"
        XLSXWriter.write(run, path, trends=None)
        wb = load_workbook(path)
        assert "Trends" not in wb.sheetnames
        wb.close()

    def test_trends_omitted_when_empty(self, tmp_path: Path) -> None:
        """Trends sheet is omitted when trends is an empty list."""
        run = _make_run_summary(
            features=[_make_feature()],
            scenarios=[_make_scenario()],
        )
        path = tmp_path / "report.xlsx"
        XLSXWriter.write(run, path, trends=[])
        wb = load_workbook(path)
        assert "Trends" not in wb.sheetnames
        wb.close()


class TestOnlyFailed:
    """only_failed filtering."""

    def test_only_failed_filters_details(self, tmp_path: Path) -> None:
        """only_failed=True filters Details but not Summary."""
        features = [_make_feature(total=2, passed=1, failed=1, pass_rate=50.0)]
        scenarios = [
            _make_scenario(scenario_name="Pass", status="passed"),
            _make_scenario(scenario_name="Fail", status="failed"),
        ]
        run = _make_run_summary(features=features, scenarios=scenarios)
        path = tmp_path / "report.xlsx"
        XLSXWriter.write(run, path, only_failed=True)
        wb = load_workbook(path)

        ws_details = wb["Details"]
        assert ws_details.max_row == 2
        assert ws_details.cell(row=2, column=2).value == "Fail"

        ws_summary = wb["Summary"]
        assert ws_summary.max_row == 2
        assert ws_summary.cell(row=2, column=1).value == "Login"
        wb.close()


class TestConditionalFormat:
    """Conditional formatting on Details Status column."""

    def test_passed_fill_green(self, tmp_path: Path) -> None:
        """Passed status has green fill."""
        run = _make_run_summary(
            features=[_make_feature()],
            scenarios=[_make_scenario(status="passed")],
        )
        path = tmp_path / "report.xlsx"
        XLSXWriter.write(run, path)
        wb = load_workbook(path)
        ws = wb["Details"]
        status_col = 3
        cell = ws.cell(row=2, column=status_col)
        assert cell.fill.start_color.rgb is not None
        assert "C6EFCE" in str(cell.fill.start_color.rgb)
        wb.close()

    def test_failed_fill_red(self, tmp_path: Path) -> None:
        """Failed status has red fill."""
        run = _make_run_summary(
            features=[_make_feature()],
            scenarios=[_make_scenario(status="failed", error_message="err")],
        )
        path = tmp_path / "report.xlsx"
        XLSXWriter.write(run, path)
        wb = load_workbook(path)
        ws = wb["Details"]
        status_col = 3
        cell = ws.cell(row=2, column=status_col)
        assert "FFC7CE" in str(cell.fill.start_color.rgb)
        wb.close()

    def test_skipped_fill_yellow(self, tmp_path: Path) -> None:
        """Skipped status has yellow fill."""
        run = _make_run_summary(
            features=[_make_feature()],
            scenarios=[_make_scenario(status="skipped")],
        )
        path = tmp_path / "report.xlsx"
        XLSXWriter.write(run, path)
        wb = load_workbook(path)
        ws = wb["Details"]
        status_col = 3
        cell = ws.cell(row=2, column=status_col)
        assert "FFEB9C" in str(cell.fill.start_color.rgb)
        wb.close()


class TestAutoFilter:
    """Auto-filter presence."""

    def test_autofilter_on_summary(self, tmp_path: Path) -> None:
        """Summary sheet has auto_filter.ref set."""
        run = _make_run_summary(
            features=[_make_feature()],
            scenarios=[_make_scenario()],
        )
        path = tmp_path / "report.xlsx"
        XLSXWriter.write(run, path)
        wb = load_workbook(path)
        ws = wb["Summary"]
        assert ws.auto_filter.ref is not None
        wb.close()

    def test_autofilter_on_details(self, tmp_path: Path) -> None:
        """Details sheet has auto_filter.ref set when there are rows."""
        run = _make_run_summary(
            features=[_make_feature()],
            scenarios=[_make_scenario()],
        )
        path = tmp_path / "report.xlsx"
        XLSXWriter.write(run, path)
        wb = load_workbook(path)
        ws = wb["Details"]
        assert ws.auto_filter.ref is not None
        wb.close()


class TestFreezePanes:
    """Freeze panes on row 1."""

    def test_freeze_panes(self, tmp_path: Path) -> None:
        """All sheets have freeze_panes set to A2."""
        run = _make_run_summary(
            features=[_make_feature()],
            scenarios=[_make_scenario()],
        )
        path = tmp_path / "report.xlsx"
        XLSXWriter.write(run, path)
        wb = load_workbook(path)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            assert ws.freeze_panes == "A2"
        wb.close()


class TestBoldHeaders:
    """Headers in bold."""

    def test_headers_bold(self, tmp_path: Path) -> None:
        """Header cells have bold font."""
        run = _make_run_summary(
            features=[_make_feature()],
            scenarios=[_make_scenario()],
        )
        path = tmp_path / "report.xlsx"
        XLSXWriter.write(run, path)
        wb = load_workbook(path)
        ws = wb["Summary"]
        assert ws.cell(row=1, column=1).font.bold is True
        ws_details = wb["Details"]
        assert ws_details.cell(row=1, column=1).font.bold is True
        wb.close()


class TestEmptyRun:
    """Empty run with no scenarios."""

    def test_empty_run(self, tmp_path: Path) -> None:
        """Empty run produces Summary with headers only, no crash."""
        run = _make_run_summary(features=[], scenarios=[])
        path = tmp_path / "report.xlsx"
        XLSXWriter.write(run, path)
        wb = load_workbook(path)
        assert "Summary" in wb.sheetnames
        ws = wb["Summary"]
        assert ws.max_row == 1
        assert ws.cell(row=1, column=1).value == "Feature"
        wb.close()


class TestDefaultSheetRemoved:
    """Default 'Sheet' is removed."""

    def test_no_default_sheet(self, tmp_path: Path) -> None:
        """The default 'Sheet' created by openpyxl is removed."""
        run = _make_run_summary(
            features=[_make_feature()],
            scenarios=[_make_scenario()],
        )
        path = tmp_path / "report.xlsx"
        XLSXWriter.write(run, path)
        wb = load_workbook(path)
        assert "Sheet" not in wb.sheetnames
        wb.close()


class TestTagsAndDuration:
    """Tags and duration serialization in Details."""

    def test_tags_semicolon_separated(self, tmp_path: Path) -> None:
        """Tags are joined with semicolons in Details."""
        run = _make_run_summary(
            features=[_make_feature()],
            scenarios=[_make_scenario(tags=["smoke", "auth"])],
        )
        path = tmp_path / "report.xlsx"
        XLSXWriter.write(run, path)
        wb = load_workbook(path)
        ws = wb["Details"]
        tags_col = 5
        assert ws.cell(row=2, column=tags_col).value == "smoke;auth"
        wb.close()

    def test_duration_formatted(self, tmp_path: Path) -> None:
        """Duration is formatted with format_duration in Details."""
        run = _make_run_summary(
            features=[_make_feature()],
            scenarios=[_make_scenario(duration=0.012)],
        )
        path = tmp_path / "report.xlsx"
        XLSXWriter.write(run, path)
        wb = load_workbook(path)
        ws = wb["Details"]
        duration_col = 4
        assert ws.cell(row=2, column=duration_col).value == "12ms"
        wb.close()


class TestEdgeCases:
    """Edge cases for branch coverage."""

    def test_undefined_status_no_fill(self, tmp_path: Path) -> None:
        """Undefined status does not get a conditional fill."""
        run = _make_run_summary(
            features=[_make_feature()],
            scenarios=[_make_scenario(status="undefined")],
        )
        path = tmp_path / "report.xlsx"
        XLSXWriter.write(run, path)
        wb = load_workbook(path)
        ws = wb["Details"]
        cell = ws.cell(row=2, column=3)
        fill = cell.fill
        assert fill.start_color.rgb is None or fill.fill_type is None
        wb.close()

    def test_columns_without_status(self, tmp_path: Path) -> None:
        """Details with columns that exclude 'status' does not crash."""
        run = _make_run_summary(
            features=[_make_feature()],
            scenarios=[_make_scenario()],
        )
        path = tmp_path / "report.xlsx"
        XLSXWriter.write(run, path, columns=["feature", "scenario"])
        wb = load_workbook(path)
        ws = wb["Details"]
        assert ws.cell(row=1, column=1).value == "feature"
        assert ws.cell(row=1, column=2).value == "scenario"
        wb.close()

    def test_is_outline_column(self, tmp_path: Path) -> None:
        """is_outline column is serialized as 'true'/'false'."""
        run = _make_run_summary(
            features=[_make_feature()],
            scenarios=[_make_scenario(is_outline=True)],
        )
        path = tmp_path / "report.xlsx"
        XLSXWriter.write(run, path, columns=["feature", "is_outline"])
        wb = load_workbook(path)
        ws = wb["Details"]
        assert ws.cell(row=2, column=2).value == "true"
        wb.close()

    def test_error_with_type_in_details(self, tmp_path: Path) -> None:
        """Error column concatenates message and type when type is present."""
        run = _make_run_summary(
            features=[_make_feature()],
            scenarios=[
                _make_scenario(
                    status="failed",
                    error_message="boom",
                    error_type="ValueError",
                ),
            ],
        )
        path = tmp_path / "report.xlsx"
        XLSXWriter.write(run, path)
        wb = load_workbook(path)
        ws = wb["Details"]
        error_col = 6
        assert ws.cell(row=2, column=error_col).value == "boom [ValueError]"
        wb.close()

    def test_only_failed_no_failed_scenarios(self, tmp_path: Path) -> None:
        """only_failed=True with no failed scenarios produces empty Details."""
        run = _make_run_summary(
            features=[_make_feature()],
            scenarios=[_make_scenario(status="passed")],
        )
        path = tmp_path / "report.xlsx"
        XLSXWriter.write(run, path, only_failed=True)
        wb = load_workbook(path)
        ws = wb["Details"]
        assert ws.max_row == 1
        wb.close()

    def test_auto_width_with_zero_values(self, tmp_path: Path) -> None:
        """Column width accounts for numeric 0 values, not treating them as empty."""
        feature = _make_feature(name="F", total=0, passed=0, failed=0, skipped=0, undefined=0)
        run = _make_run_summary(features=[feature])
        path = tmp_path / "report.xlsx"
        XLSXWriter.write(run, path)
        wb = load_workbook(path)
        ws = wb["Summary"]
        total_col_letter = "B"
        width = ws.column_dimensions[total_col_letter].width
        assert width >= 3
        wb.close()
