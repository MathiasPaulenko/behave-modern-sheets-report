"""XLSX writing logic using openpyxl.

This module provides ``XLSXWriter``, a static utility that serialises a
``RunSummary`` into a multi-sheet Excel workbook.  Requires the ``openpyxl``
package (installed via the ``[xlsx]`` extra).
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.styles import Font, PatternFill  # type: ignore[import-untyped]
from openpyxl.worksheet.worksheet import Worksheet  # type: ignore[import-untyped]

from .csv_writer import scenario_cell_value
from .models import FeatureSummary, HistoryEntry, RunSummary, ScenarioResult
from .utils import (
    DEFAULT_COLUMNS,
    STATUS_FAILED,
    STATUS_PASSED,
    STATUS_SKIPPED,
    format_duration,
)

_FILL_PASSED = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
_FILL_FAILED = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
_FILL_SKIPPED = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")

_BOLD_FONT = Font(bold=True)

_SUMMARY_HEADERS = [
    "Feature",
    "Total",
    "Passed",
    "Failed",
    "Skipped",
    "Undefined",
    "Pass Rate",
    "Duration",
]

_FAILURES_HEADERS = [
    "Feature",
    "Scenario",
    "Error",
    "Error Type",
    "Traceback",
    "File",
    "Line",
]

_TRENDS_HEADERS = [
    "Timestamp",
    "Pass Rate",
    "Total Features",
    "Total Scenarios",
    "Passed",
    "Failed",
    "Skipped",
    "Undefined",
    "Duration",
]

__all__ = ["XLSXWriter"]


class XLSXWriter:
    """Write scenario results as a multi-sheet XLSX workbook."""

    @staticmethod
    def write(
        run_summary: RunSummary,
        path: Path,
        columns: list[str] | None = None,
        only_failed: bool = False,
        trends: list[HistoryEntry] | None = None,
    ) -> None:
        """Serialize ``run_summary`` into an XLSX file at ``path``.

        Args:
            run_summary: The complete run summary containing features and
                scenarios to write.
            path: Destination file path for the ``.xlsx`` file.
            columns: Ordered list of column names for the Details sheet.
                When ``None``, ``DEFAULT_COLUMNS`` is used.
            only_failed: When ``True``, the Details sheet only includes
                failed scenarios.  Summary and Trends remain complete.
            trends: Historical run entries for the Trends sheet.  When
                ``None`` or empty, the Trends sheet is omitted.
        """
        cols = columns if columns is not None else DEFAULT_COLUMNS
        wb = Workbook()

        XLSXWriter._write_summary(wb, run_summary.features)
        XLSXWriter._write_details(wb, run_summary.scenarios, cols, only_failed)
        XLSXWriter._write_failures(wb, run_summary.scenarios)
        if trends is not None and len(trends) > 0:
            XLSXWriter._write_trends(wb, trends)

        default_sheet = wb["Sheet"]
        wb.remove(default_sheet)

        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)

    @staticmethod
    def _write_summary(wb: Workbook, features: list[FeatureSummary]) -> None:
        """Write the Summary sheet with one row per feature.

        Args:
            wb: The workbook to add the sheet to.
            features: List of feature summaries to write.
        """
        ws: Worksheet = wb.create_sheet("Summary")
        ws.append(_SUMMARY_HEADERS)
        XLSXWriter._style_header(ws, len(_SUMMARY_HEADERS))

        for feature in features:
            ws.append([
                feature.feature_name,
                feature.total_scenarios,
                feature.passed,
                feature.failed,
                feature.skipped,
                feature.undefined,
                f"{feature.pass_rate:.1f}%",
                format_duration(feature.duration),
            ])

        if len(features) > 0:
            ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = "A2"
        XLSXWriter._auto_width(ws)

    @staticmethod
    def _write_details(
        wb: Workbook,
        scenarios: list[ScenarioResult],
        columns: list[str],
        only_failed: bool,
    ) -> None:
        """Write the Details sheet with one row per scenario.

        Args:
            wb: The workbook to add the sheet to.
            scenarios: List of scenario results to write.
            columns: Column names to include.
            only_failed: When ``True``, only failed scenarios are included.
        """
        ws: Worksheet = wb.create_sheet("Details")
        ws.append(columns)
        XLSXWriter._style_header(ws, len(columns))

        status_col = columns.index("status") + 1 if "status" in columns else 0

        row_idx = 2
        for scenario in scenarios:
            if only_failed and scenario.status != STATUS_FAILED:
                continue
            row = XLSXWriter._scenario_row(scenario, columns)
            ws.append(row)

            if status_col > 0:
                fill = XLSXWriter._status_fill(scenario.status)
                if fill is not None:
                    ws.cell(row=row_idx, column=status_col).fill = fill
            row_idx += 1

        data_rows = row_idx - 2
        if data_rows > 0:
            ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = "A2"
        XLSXWriter._auto_width(ws)

    @staticmethod
    def _write_failures(wb: Workbook, scenarios: list[ScenarioResult]) -> None:
        """Write the Failures sheet with only failed scenarios.

        Args:
            wb: The workbook to add the sheet to.
            scenarios: List of scenario results to filter.
        """
        ws: Worksheet = wb.create_sheet("Failures")
        ws.append(_FAILURES_HEADERS)
        XLSXWriter._style_header(ws, len(_FAILURES_HEADERS))

        for scenario in scenarios:
            if scenario.status != STATUS_FAILED:
                continue
            ws.append([
                scenario.feature_name,
                scenario.scenario_name,
                scenario.error_message,
                scenario.error_type,
                scenario.traceback,
                scenario.file,
                scenario.line,
            ])

        failed_count = sum(1 for s in scenarios if s.status == STATUS_FAILED)
        if failed_count > 0:
            ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = "A2"
        XLSXWriter._auto_width(ws)

    @staticmethod
    def _write_trends(wb: Workbook, trends: list[HistoryEntry]) -> None:
        """Write the Trends sheet with one row per history entry.

        Args:
            wb: The workbook to add the sheet to.
            trends: List of historical run entries.
        """
        ws: Worksheet = wb.create_sheet("Trends")
        ws.append(_TRENDS_HEADERS)
        XLSXWriter._style_header(ws, len(_TRENDS_HEADERS))

        for entry in trends:
            ws.append([
                entry.timestamp,
                f"{entry.pass_rate:.1f}%",
                entry.total_features,
                entry.total_scenarios,
                entry.passed,
                entry.failed,
                entry.skipped,
                entry.undefined,
                format_duration(entry.duration),
            ])

        ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = "A2"
        XLSXWriter._auto_width(ws)

    @staticmethod
    def _scenario_row(scenario: ScenarioResult, columns: list[str]) -> list[str]:
        """Build a single row of cell values from ``scenario``.

        Args:
            scenario: The scenario result to serialize.
            columns: The column names to include.

        Returns:
            A list of string values, one per column.
        """
        return [scenario_cell_value(scenario, col) for col in columns]

    @staticmethod
    def _status_fill(status: str) -> PatternFill | None:
        """Return the conditional fill for a scenario status.

        Args:
            status: The canonical status string.

        Returns:
            A ``PatternFill`` for known statuses, or ``None``.
        """
        if status == STATUS_PASSED:
            return _FILL_PASSED
        if status == STATUS_FAILED:
            return _FILL_FAILED
        if status == STATUS_SKIPPED:
            return _FILL_SKIPPED
        return None

    @staticmethod
    def _style_header(ws: Worksheet, num_cols: int) -> None:
        """Apply bold font to the header row.

        Args:
            ws: The worksheet to style.
            num_cols: Number of columns in the header.
        """
        for col_idx in range(1, num_cols + 1):
            ws.cell(row=1, column=col_idx).font = _BOLD_FONT

    @staticmethod
    def _auto_width(ws: Worksheet) -> None:
        """Set column widths based on maximum content length.

        Args:
            ws: The worksheet to adjust.
        """
        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                value = cell.value
                length = len(str(value)) if value is not None else 0
                if length > max_len:
                    max_len = length
            ws.column_dimensions[col_letter].width = max_len + 2
